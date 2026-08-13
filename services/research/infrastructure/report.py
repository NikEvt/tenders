"""Отчёт по исследованию в xlsx.

Два свойства, оба выведены из того, как отчёт однажды уронил прогон.

**Сборка отчёта не имеет права оборвать основную работу.** В прогоне ХПК/БПК
она вызывалась прямо из цикла разбора, упала на недопустимом символе и унесла с
собой полчаса работы. Данные лежат в базе, отчёт из них пересобирается когда
угодно — значит, его сбой обязан стоить строки в логе. Отсюда `safe_report`.

**Файл подменяется целиком и атомарно.** Отчёт переписывается по ходу прогона,
и открытый в этот момент файл не должен оказаться обрезанным: запись идёт во
временный файл рядом, затем `os.replace`. В пределах одной файловой системы это
атомарная операция.
"""

from __future__ import annotations

import os
import tempfile
from collections.abc import Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from libs.shared.logging import get_logger
from libs.shared.text_sanitize import cell
from services.research.application.ports import TenderVerdict
from services.research.domain.criteria import Confidence
from services.research.domain.market import MarketSummary, MarketTender, summarize

log = get_logger(__name__)

NOTICE_URL = "https://zakupki.gov.ru/epz/order/notice/ea44/view/common-info.html?regNumber={}"

SUMMARY_HEADERS = [
    "Ссылка на закупку",
    "Рег. номер",
    "Наименование",
    "НМЦК, руб.",
    "Заказчик",
    "ИНН",
    "Регион",
    "ОКПД2",
    "Решение",
    "Кем решено",
    "Основание",
    "Документ-источник",
    "Цитата",
]

MARKET_HEADERS = ["Разрез", "Ключ", "Значение", "Закупок", "Сумма, руб.", "Средний чек"]


@dataclass(frozen=True, slots=True)
class ReportRow:
    """Строка отчёта: закупка вместе с решением и цитатой."""

    verdict: TenderVerdict
    tender: MarketTender
    quote: str = ""
    file_name: str | None = None


def build_report(path: Path, rows: Sequence[ReportRow]) -> None:
    """Собирает отчёт и атомарно подменяет им прежний файл."""
    import openpyxl

    workbook = openpyxl.Workbook()
    confirmed_sheet = workbook.active
    confirmed_sheet.title = "Подтверждённые"
    rejected_sheet = workbook.create_sheet("Отклонённые")
    market_sheet = workbook.create_sheet("Рынок")

    for sheet in (confirmed_sheet, rejected_sheet):
        sheet.append(SUMMARY_HEADERS)
    market_sheet.append(MARKET_HEADERS)

    confirmed: list[MarketTender] = []
    for row in rows:
        sheet = (
            confirmed_sheet
            if row.verdict.confidence is Confidence.CONFIRMED
            else rejected_sheet
        )
        sheet.append([cell(value) for value in _line(row)])
        if row.verdict.confidence is Confidence.CONFIRMED:
            confirmed.append(row.tender)

    for line in _market_lines(summarize(confirmed)):
        market_sheet.append([cell(value) for value in line])

    _save_atomically(workbook, path)


def _line(row: ReportRow) -> list[object]:
    tender = row.tender
    return [
        NOTICE_URL.format(tender.reg_num),
        tender.reg_num,
        tender.name,
        float(tender.price) if tender.price is not None else None,
        tender.customer_name,
        tender.customer_inn,
        tender.region_code,
        tender.okpd2_code,
        str(row.verdict.confidence),
        row.verdict.decided_by,
        row.verdict.reason,
        row.file_name,
        row.quote,
    ]


def _market_lines(summary: MarketSummary) -> list[list[object]]:
    lines: list[list[object]] = [
        ["итого", "", "закупок", summary.total_count, float(summary.total_value), None],
        [
            "итого",
            "",
            "медиана / среднее",
            summary.priced_count,
            float(summary.median_price or 0),
            float(summary.average_price or 0),
        ],
    ]
    for name, buckets in (
        ("регион", summary.by_region),
        ("заказчик", summary.by_customer),
        ("ОКПД2", summary.by_okpd2),
    ):
        for bucket in buckets:
            lines.append(
                [
                    name,
                    bucket.key,
                    bucket.label,
                    bucket.count,
                    float(bucket.total),
                    float(bucket.average) if bucket.average is not None else None,
                ]
            )
    return lines


def _save_atomically(workbook: Any, path: Path) -> None:
    """Пишет рядом и подменяет одним движением.

    Открытый в этот момент отчёт не должен оказаться обрезанным: пока идёт
    запись, прежний файл цел, а `os.replace` в пределах одной файловой системы
    атомарен.
    """
    path.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary = tempfile.mkstemp(dir=str(path.parent), suffix=".xlsx")
    os.close(handle)
    try:
        workbook.save(temporary)
        os.replace(temporary, path)
    except BaseException:
        Path(temporary).unlink(missing_ok=True)
        raise


def safe_report(path: Path, rows: Sequence[ReportRow]) -> bool:
    """Собирает отчёт, но не даёт ему уронить прогон.

    Вспомогательное действие не имеет права стоить основного. Однажды уже
    стоило: сборка отчёта упала на символе из OCR и унесла с собой полчаса
    разбора. Данные в базе, отчёт пересобирается когда угодно.
    """
    try:
        build_report(path, rows)
    except Exception as exc:
        log.warning("research.report_failed", path=str(path), error=str(exc))
        return False
    return True
