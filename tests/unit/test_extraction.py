"""Извлечение текста из форматов, реально встречающихся во вложениях ЕИС."""

from __future__ import annotations

import io
import zipfile

import pytest

from services.docs_worker.domain.models import ExtractedPage, ExtractedText
from services.docs_worker.infrastructure.extraction.archives import (
    MAX_ENTRIES,
    ZipExtractor,
)
from services.docs_worker.infrastructure.extraction.base import (
    ExtractorRegistry,
    SkippedExtractor,
)
from services.docs_worker.infrastructure.extraction.documents import (
    DocxExtractor,
    PdfExtractor,
    PlainTextExtractor,
    XlsxExtractor,
    normalize,
)
from services.docs_worker.infrastructure.extraction.ocr import PdfWithOcrFallback


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    import docx

    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        created = document.add_table(rows=len(table), cols=len(table[0]))
        for row_index, row in enumerate(table):
            for cell_index, value in enumerate(row):
                created.cell(row_index, cell_index).text = value

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_xlsx(sheets: dict[str, list[list[object]]]) -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for title, rows in sheets.items():
        sheet = workbook.create_sheet(title)
        for row in rows:
            sheet.append(row)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def make_zip(entries: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in entries.items():
            archive.writestr(name, content)
    return buffer.getvalue()


class TestDocx:
    def test_paragraphs_and_tables(self) -> None:
        content = make_docx(
            ["Техническое задание", "Поставка газа в баллонах"],
            table=[["Наименование", "Количество"], ["Кислород технический", "40 баллонов"]],
        )
        result = DocxExtractor().extract(content, "tz.docx")

        assert "Техническое задание" in result.content
        # Требования в ТЗ чаще всего лежат именно в таблицах.
        assert "Кислород технический" in result.content
        assert "40 баллонов" in result.content
        assert result.extractor == "docx"
        assert not result.ocr_used

    def test_empty_document_is_not_an_error(self) -> None:
        result = DocxExtractor().extract(make_docx([]), "empty.docx")
        assert result.is_empty
        assert result.pages == []


class TestXlsx:
    def test_each_sheet_becomes_a_page(self) -> None:
        content = make_xlsx(
            {
                "Смета": [["Позиция", "Цена"], ["Лампа E27", 1200]],
                "Расчёт НМЦК": [["Метод", "Сопоставимых цен"]],
            }
        )
        result = XlsxExtractor().extract(content, "raschet.xlsx")

        assert result.page_count == 2
        assert "Лампа E27" in result.content
        assert "Сопоставимых цен" in result.content
        assert "Лист: Смета" in result.pages[0].text


class TestPlainText:
    @pytest.mark.parametrize("encoding", ["utf-8", "cp1251"])
    def test_russian_encodings(self, encoding: str) -> None:
        result = PlainTextExtractor().extract(
            "Поставка кислорода".encode(encoding), "note.txt"
        )
        assert "Поставка кислорода" in result.content

    def test_html_tags_are_stripped(self) -> None:
        result = PlainTextExtractor().extract(
            b"<html><body><p>\xd0\xa2\xd0\x97</p></body></html>", "doc.html"
        )
        assert "<" not in result.content
        assert "ТЗ" in result.content


class TestZip:
    def test_returns_embedded_files_not_text(self) -> None:
        content = make_zip(
            {
                "ТЗ.txt": "Поставка газа".encode(),
                "смета.txt": "1200 рублей".encode(),
                "подпись.sig": b"\x00\x01",
                "dir/": b"",
            }
        )
        result = ZipExtractor().extract(content, "komplekt.zip")

        names = {f.file_name for f in result.embedded_files}
        assert names == {"ТЗ.txt", "смета.txt"}, "подписи и каталоги пропускаются"
        assert result.is_empty, "архив сам по себе текста не содержит"

    def test_path_traversal_entries_are_dropped(self) -> None:
        content = make_zip({"../../etc/passwd": b"root", "ok.txt": b"data"})
        result = ZipExtractor().extract(content, "evil.zip")
        assert {f.file_name for f in result.embedded_files} == {"ok.txt"}

    def test_entry_count_is_capped(self) -> None:
        content = make_zip({f"f{i}.txt": b"x" for i in range(MAX_ENTRIES + 50)})
        result = ZipExtractor().extract(content, "many.zip")
        assert len(result.embedded_files) <= MAX_ENTRIES


class TestRegistry:
    @pytest.fixture
    def registry(self) -> ExtractorRegistry:
        return ExtractorRegistry(
            [SkippedExtractor(), DocxExtractor(), XlsxExtractor(), PlainTextExtractor()]
        )

    def test_signatures_are_skipped(self, registry: ExtractorRegistry) -> None:
        result = registry.extract(b"binary-signature", "contract.pdf.sig")
        assert result.extractor == "skipped"

    def test_unknown_extension_is_reported_not_raised(
        self, registry: ExtractorRegistry
    ) -> None:
        result = registry.extract(b"\x00\x01", "drawing.dwg")
        assert result.extractor == "unsupported"
        assert result.is_empty

    def test_broken_file_does_not_propagate(self, registry: ExtractorRegistry) -> None:
        """Один нечитаемый файл не должен ронять обработку остальных вложений."""
        result = registry.extract(b"not really a docx", "broken.docx")
        assert result.extractor == "failed"
        assert result.is_empty

    def test_oversized_file_is_refused(self, registry: ExtractorRegistry) -> None:
        from services.docs_worker.infrastructure.extraction.base import MAX_EXTRACTABLE_BYTES

        result = registry.extract(b"x" * (MAX_EXTRACTABLE_BYTES + 1), "huge.txt")
        assert result.extractor == "too_large"


class TestOcrFallback:
    class FakeText:
        """Экстрактор текстового слоя, отдающий заданный результат."""

        extensions = frozenset({"pdf"})

        def __init__(self, result: ExtractedText) -> None:
            self._result = result

        def supports(self, file_name: str, content_type: str | None) -> bool:
            return file_name.endswith(".pdf")

        def extract(self, content: bytes, file_name: str) -> ExtractedText:
            return self._result

    def test_rich_text_layer_skips_ocr(self) -> None:
        rich = ExtractedText(
            pages=[ExtractedPage(number=1, text="а" * 500)], extractor="pdfium"
        )
        ocr = self.FakeText(ExtractedText(pages=[], extractor="tesseract", ocr_used=True))

        result = PdfWithOcrFallback(self.FakeText(rich), ocr).extract(b"", "doc.pdf")
        assert result.extractor == "pdfium"
        assert not result.ocr_used

    def test_scan_falls_back_to_ocr(self) -> None:
        # Пустая шапка вместо текста — типичный скан ТЗ.
        scan = ExtractedText(pages=[ExtractedPage(number=1, text="стр. 1")], extractor="pdfium")
        recognized = ExtractedText(
            pages=[ExtractedPage(number=1, text="Поставка газа в баллонах " * 20)],
            extractor="tesseract",
            ocr_used=True,
        )

        result = PdfWithOcrFallback(self.FakeText(scan), self.FakeText(recognized)).extract(
            b"", "scan.pdf"
        )
        assert result.ocr_used
        assert "Поставка газа" in result.content

    def test_ocr_worse_than_text_layer_is_discarded(self) -> None:
        text_layer = ExtractedText(
            pages=[ExtractedPage(number=1, text="Короткий, но осмысленный текст")],
            extractor="pdfium",
        )
        garbage = ExtractedText(
            pages=[ExtractedPage(number=1, text="~")], extractor="tesseract", ocr_used=True
        )

        result = PdfWithOcrFallback(
            self.FakeText(text_layer), self.FakeText(garbage)
        ).extract(b"", "doc.pdf")
        assert result.extractor == "pdfium"


def test_normalize_collapses_whitespace_but_keeps_paragraphs() -> None:
    assert normalize("Поставка   газа\r\n\n\n\nв  баллонах") == "Поставка газа\n\nв баллонах"


def test_pdf_extractor_reads_generated_pdf() -> None:
    """Проверка на настоящем PDF, а не на моке."""
    pypdfium2 = pytest.importorskip("pypdfium2")

    document = pypdfium2.PdfDocument.new()
    document.new_page(595, 842)
    buffer = io.BytesIO()
    document.save(buffer)
    document.close()

    result = PdfExtractor().extract(buffer.getvalue(), "blank.pdf")
    assert result.extractor == "pdfium"
    assert result.page_count == 1
    # Пустая страница — валидный результат, а не ошибка (LSP).
    assert result.is_empty
    assert result.looks_like_scan()
