"""Пул процессов извлечения: тот же результат, но переживающий крах разборщика.

Главное, ради чего пул существует, проверяется не здесь, а замером
(`scripts/bench_extraction.py`): на потоках тот же реестр экстракторов роняет
процесс на двух параллельных PDF. Воспроизводить нативный сегфолт в тесте —
значит получить тест, который иногда падает сам; вместо этого проверяется то,
что действительно можно проверить надёжно: контракт порта, живучесть при
падении процесса и смена размера пула.
"""

from __future__ import annotations

import io
import os
from concurrent.futures import ThreadPoolExecutor
from concurrent.futures.process import BrokenProcessPool

import pytest

from services.docs_worker.bootstrap import build_extraction
from services.docs_worker.infrastructure.extraction.pool import ProcessPoolExtraction


def make_docx(text: str) -> bytes:
    import docx

    document = docx.Document()
    document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_xlsx(rows: list[list[object]]) -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.fixture(scope="module")
def pool():
    extraction = ProcessPoolExtraction(workers=2)
    yield extraction
    extraction.shutdown()


class TestPortContract:
    """Пул обязан быть неотличим от разбора на месте (LSP)."""

    def test_matches_in_process_extraction(self, pool: ProcessPoolExtraction) -> None:
        content = make_docx("Определение ХПК в сточных водах по ПНД Ф 14.1:2.100-97")

        in_process = build_extraction().extract(content, "тз.docx", None)
        in_pool = pool.extract(content, "тз.docx", None)

        assert in_pool.extractor == in_process.extractor
        assert in_pool.char_count == in_process.char_count
        assert in_pool.content == in_process.content

    def test_carries_pages_and_metadata_back(self, pool: ProcessPoolExtraction) -> None:
        """Результат едет через pickle — поля не должны потеряться по дороге."""
        result = pool.extract(make_xlsx([["БПК5", 12.5], ["ХПК", 800]]), "смета.xlsx", None)

        assert result.extractor
        assert result.pages
        assert "БПК5" in result.content
        assert result.ocr_used is False

    def test_empty_result_is_not_an_error(self, pool: ProcessPoolExtraction) -> None:
        """Пустой текст — валидный результат, как и при разборе на месте."""
        result = pool.extract(b"", "пустой.txt", None)
        assert result.is_empty

    def test_unknown_format_does_not_raise(self, pool: ProcessPoolExtraction) -> None:
        result = pool.extract(b"\x00\x01\x02", "подпись.sig", None)
        assert result.is_empty


class TestParallelWork:
    def test_concurrent_calls_from_threads(self, pool: ProcessPoolExtraction) -> None:
        """Ровно та нагрузка, на которой пул потоков умирает нативным сигналом.

        Вызовы идут из нескольких потоков — так их и делает сценарий через
        `asyncio.to_thread`. Разбор при этом происходит в чужих процессах,
        поэтому делить между потоками нечего.
        """
        documents = [make_docx(f"Протокол {i}: БПК5 и ХПК") for i in range(12)]

        with ThreadPoolExecutor(6) as threads:
            results = list(
                threads.map(
                    lambda item: pool.extract(item[1], f"д-{item[0]}.docx", None),
                    enumerate(documents),
                )
            )

        assert len(results) == 12
        assert all("БПК5" in result.content for result in results)


def _suicide(_content: bytes, _name: str, _type: str | None):
    """Убивает рабочий процесс так же, как это делает битый PDF."""
    os._exit(1)


def _read_omp_threads(*_args: object) -> str | None:
    """Читает переменную внутри рабочего процесса.

    Функция уровня модуля, потому что задача едет в процесс через pickle:
    `os.environ.get` — связанный метод объекта, собранного в замыкании
    `_createenviron`, и не сериализуется.
    """
    return os.environ.get("OMP_NUM_THREADS")


class TestPoolSurvivesACrash:
    def test_rebuilds_after_a_worker_dies(self, monkeypatch) -> None:
        """Развал пула не должен делать сервис непригодным навсегда."""
        from services.docs_worker.infrastructure.extraction import pool as module

        extraction = ProcessPoolExtraction(workers=2)
        try:
            monkeypatch.setattr(module, "_extract_in_worker", _suicide)
            with pytest.raises(BrokenProcessPool):
                extraction.extract(b"x", "убийца.pdf", None)

            # Пул пересобран — следующий документ разбирается как ни в чём не бывало.
            monkeypatch.undo()
            result = extraction.extract(make_docx("после сбоя"), "живой.docx", None)
            assert "после сбоя" in result.content
        finally:
            extraction.shutdown()

    def test_a_killer_file_does_not_take_its_neighbours(self, monkeypatch) -> None:
        """Сбойным помечается один документ, остальные обязаны дойти."""
        from services.docs_worker.infrastructure.extraction import pool as module

        extraction = ProcessPoolExtraction(workers=2)
        try:
            monkeypatch.setattr(module, "_extract_in_worker", _suicide)
            with pytest.raises(BrokenProcessPool):
                extraction.extract(b"x", "убийца.pdf", None)
            monkeypatch.undo()

            survivors = [
                extraction.extract(make_docx(f"документ {i}"), f"д-{i}.docx", None)
                for i in range(3)
            ]
            assert all(f"документ {i}" in survivors[i].content for i in range(3))
        finally:
            extraction.shutdown()


class TestResizing:
    def test_resize_changes_the_worker_count(self) -> None:
        extraction = ProcessPoolExtraction(workers=2)
        try:
            extraction.extract(make_docx("до"), "до.docx", None)
            extraction.resize(4)
            assert extraction.workers == 4

            result = extraction.extract(make_docx("после"), "после.docx", None)
            assert "после" in result.content
        finally:
            extraction.shutdown()

    def test_resize_to_the_same_size_keeps_the_pool(self) -> None:
        """Пересборка на ровном месте стоила бы простоя на каждом опросе."""
        extraction = ProcessPoolExtraction(workers=2)
        try:
            extraction.extract(make_docx("раз"), "раз.docx", None)
            before = extraction._generation
            extraction.resize(2)
            extraction.extract(make_docx("два"), "два.docx", None)
            assert extraction._generation == before
        finally:
            extraction.shutdown()

    def test_never_drops_below_one_worker(self) -> None:
        extraction = ProcessPoolExtraction(workers=4)
        try:
            extraction.resize(0)
            assert extraction.workers == 1
        finally:
            extraction.shutdown()


class TestWorkerEnvironment:
    def test_openmp_is_pinned_to_one_thread(self) -> None:
        """N процессов × N потоков поделили бы то же железо дважды."""
        extraction = ProcessPoolExtraction(workers=1)
        try:
            from services.docs_worker.infrastructure.extraction import pool as module

            pool_obj, _ = extraction._ensure_pool()
            assert pool_obj.submit(_read_omp_threads).result() == "1"
            assert module.MAX_TASKS_PER_CHILD > 0
        finally:
            extraction.shutdown()
