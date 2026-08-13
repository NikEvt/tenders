"""Сборка эталонных наборов из результатов прогона ХПК/БПК 11–12 августа 2026.

Прогон оставил после себя размеченный вручную результат: 56 подтверждённых
закупок, 20 отклонённых с причинами и 333 находки с настоящими цитатами. Это
единственный в проекте набор, где известен правильный ответ, поэтому он
становится оракулом приёмки нового движка отбора.

Скрипт разовый по смыслу, но не по исполнению: фикстуры пересобираются из тех же
источников и обязаны получаться байт-в-байт теми же. Отсюда сортировки на каждом
шаге и фиксированное зерно выборки — иначе диff фикстуры шумит на ровном месте и
перестаёт что-либо значить.

    python scripts/build_golden_fixtures.py

Источники (не входят в репозиторий, лежат локально после прогона):
    data/market_analysis/hpk.sqlite
    data/market_analysis/hpk_bpk_подтверждённые.xlsx
"""

from __future__ import annotations

import json
import random
import sqlite3
from pathlib import Path
from typing import Any

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
SQLITE = ROOT / "data/market_analysis/hpk.sqlite"
XLSX = ROOT / "data/market_analysis/hpk_bpk_подтверждённые.xlsx"
OUT = ROOT / "tests/fixtures"

#: Зерно выборки. Меняется только вместе с осознанным решением пересобрать
#: выборку заново — иначе фикстура перестаёт быть сравнимой с прежней.
SEED = 20260812

#: Причина отклонения → класс ложного срабатывания. Ключ ищется как префикс
#: причины, записанной вручную при разборе результата.
#:
#: «ХПК Мариинского театра» — художественно-производственный комбинат,
#: «БПК СИЗО» — банно-прачечный. Обе аббревиатуры совпадают с искомыми буква в
#: букву, и различает их только окружение, а не шаблон.
REJECT_CLASSES: tuple[tuple[str, str], ...] = (
    ("медицина", "medicine"),
    ("микробиология", "medicine"),
    ("артефакт шаблона", "template_artifact"),
    ("маркировка изделия", "product_marking"),
    ("банно-прачечный комбинат", "object_name"),
    ("ХПК Мариинского театра", "object_name"),
    ("здание", "object_name"),
)


def reject_class(reason: str) -> str:
    for prefix, name in REJECT_CLASSES:
        if reason.startswith(prefix):
            return name
    raise SystemExit(f"неизвестный класс отклонения: {reason!r}")


def _split_codes(raw: str | None) -> list[str]:
    return [part.strip() for part in (raw or "").split(",") if part.strip()]


def load_labels() -> dict[str, dict[str, Any]]:
    """Рег. номер → вердикт человека."""
    book = openpyxl.load_workbook(XLSX, read_only=True)

    labels: dict[str, dict[str, Any]] = {}
    for row in book["Подтверждённые"].iter_rows(min_row=2, values_only=True):
        if row[1]:
            labels[str(row[1])] = {"label": "confirmed", "reject_class": None, "reason": None}

    for row in book["Отклонённые"].iter_rows(min_row=2, values_only=True):
        if not row[0]:
            continue
        reason = str(row[4])
        labels[str(row[0])] = {
            "label": "rejected",
            "reject_class": reject_class(reason),
            "reason": reason,
        }
    return labels


def build_golden(db: sqlite3.Connection, labels: dict[str, dict[str, Any]]) -> dict[str, Any]:
    """76 размеченных закупок вместе с настоящими цитатами из прогона."""
    hits_by_tender: dict[str, list[dict[str, Any]]] = {}
    for row in db.execute(
        "SELECT reg_num, source, term, quote, file_name, page FROM hits ORDER BY id"
    ):
        hits_by_tender.setdefault(row["reg_num"], []).append(
            {
                "source": row["source"],
                "term": row["term"],
                "quote": row["quote"],
                "file_name": row["file_name"],
                "page": row["page"],
            }
        )

    tenders = []
    for reg_num in sorted(labels):
        row = db.execute(
            "SELECT reg_num, region, name, description, price, customer_name,"
            "       customer_inn, okpd2_codes, okpd2_names, publish_date, end_date"
            "  FROM tenders WHERE reg_num = ?",
            (reg_num,),
        ).fetchone()
        if row is None:
            raise SystemExit(f"закупка {reg_num} размечена, но её нет в базе прогона")

        verdict = labels[reg_num]
        tenders.append(
            {
                "reg_num": row["reg_num"],
                "region": row["region"],
                "name": row["name"],
                "description": row["description"],
                "price": row["price"],
                "customer_name": row["customer_name"],
                "customer_inn": row["customer_inn"],
                "okpd2_codes": _split_codes(row["okpd2_codes"]),
                "okpd2_names": _split_codes(row["okpd2_names"]),
                "publish_date": row["publish_date"],
                "end_date": row["end_date"],
                "label": verdict["label"],
                "reject_class": verdict["reject_class"],
                "reject_reason": verdict["reason"],
                "hits": hits_by_tender.get(reg_num, []),
            }
        )

    confirmed = sum(1 for t in tenders if t["label"] == "confirmed")
    return {
        "meta": {
            "source": "прогон ХПК/БПК 11–12 августа 2026",
            "confirmed": confirmed,
            "rejected": len(tenders) - confirmed,
            "note": (
                "Разметка сделана человеком по цитатам. Полнота движка меряется по "
                "confirmed, точность — по rejected с разбивкой на reject_class."
            ),
        },
        "tenders": tenders,
    }


def build_prefilter_sample(db: sqlite3.Connection, size: int = 600) -> dict[str, Any]:
    """Выборка карточек с ответом действующего предфильтра.

    Характеризует поведение, а не задаёт желаемое: предфильтр намеренно широкий
    и переносится в новый движок как есть. Тест на этой выборке ловит случайное
    расхождение при переписывании, а не оценивает качество отбора.
    """
    rows = db.execute(
        "SELECT reg_num, name, description, okpd2_codes, okpd2_names, is_candidate, card_hit"
        "  FROM tenders ORDER BY reg_num"
    ).fetchall()

    candidates = [r for r in rows if r["is_candidate"]]
    rest = [r for r in rows if not r["is_candidate"]]

    rng = random.Random(SEED)
    half = size // 2
    picked = rng.sample(candidates, min(half, len(candidates)))
    picked += rng.sample(rest, min(size - half, len(rest)))

    return {
        "meta": {
            "note": "Ответ предфильтра по карточке на день прогона. Характеризация.",
            "candidates": sum(1 for r in picked if r["is_candidate"]),
            "total": len(picked),
        },
        "tenders": sorted(
            (
                {
                    "reg_num": r["reg_num"],
                    "name": r["name"],
                    "description": r["description"],
                    "okpd2_codes": _split_codes(r["okpd2_codes"]),
                    "okpd2_names": _split_codes(r["okpd2_names"]),
                    "is_candidate": bool(r["is_candidate"]),
                }
                for r in picked
            ),
            key=lambda t: t["reg_num"],
        ),
    }


def build_admission_sample(db: sqlite3.Connection, per_priority: int = 60) -> dict[str, Any]:
    """Вложения с приоритетом, присвоенным в прогоне, по всем классам приоритета."""
    rng = random.Random(SEED)
    picked = []
    for priority in (0, 1, 2, 3, 9):
        rows = db.execute(
            "SELECT reg_num, file_name, doc_kind_name, file_size, priority"
            "  FROM attachments WHERE priority = ? ORDER BY id",
            (priority,),
        ).fetchall()
        picked += rng.sample(rows, min(per_priority, len(rows)))

    return {
        "meta": {
            "note": (
                "Приоритет документа по имени и виду, как его считал прогон. "
                "Порядок разбора: ТЗ → обоснование НМЦК → приложения → контракты → прочее."
            ),
            "total": len(picked),
        },
        "attachments": sorted(
            (
                {
                    "file_name": r["file_name"],
                    "doc_kind_name": r["doc_kind_name"],
                    "file_size": r["file_size"],
                    "priority": r["priority"],
                }
                for r in picked
            ),
            key=lambda a: (a["priority"], a["file_name"] or "", a["file_size"] or 0),
        ),
    }


def build_funnel(db: sqlite3.Connection) -> dict[str, Any]:
    """Опорные числа воронки прогона — знаменатели для сверки на этапе 6.

    Ретроспектива отдельно предупреждает: «ноль находок» без знаменателя не
    значит ничего. Поэтому вместе с находками сохраняются и величины, из
    которых они получены.
    """
    scalar = lambda sql: db.execute(sql).fetchone()[0]  # noqa: E731
    by_status = {
        row["status"]: row["n"]
        for row in db.execute(
            "SELECT status, COUNT(*) AS n FROM attachments GROUP BY status ORDER BY status"
        )
    }
    return {
        "tenders_total": scalar("SELECT COUNT(*) FROM tenders"),
        "tenders_candidate": scalar("SELECT COUNT(*) FROM tenders WHERE is_candidate = 1"),
        "attachments_total": scalar("SELECT COUNT(*) FROM attachments"),
        "attachments_by_status": by_status,
        "hits_total": scalar("SELECT COUNT(*) FROM hits"),
        "tenders_with_hits": scalar("SELECT COUNT(DISTINCT reg_num) FROM hits"),
        "crawl_days": scalar("SELECT COUNT(*) FROM crawl_state WHERE status = 'success'"),
    }


def write(name: str, payload: dict[str, Any]) -> None:
    path = OUT / name
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=False) + "\n",
        encoding="utf-8",
    )
    print(f"{path.relative_to(ROOT)}: {path.stat().st_size / 1024:.0f} КБ")


def main() -> None:
    for source in (SQLITE, XLSX):
        if not source.exists():
            raise SystemExit(f"нет источника: {source}")

    OUT.mkdir(parents=True, exist_ok=True)

    # Только чтение: 285 МБ состояния прогона правиться не должны ни при каких
    # обстоятельствах — второго такого набора не будет.
    db = sqlite3.connect(f"file:{SQLITE}?mode=ro", uri=True)
    db.row_factory = sqlite3.Row
    try:
        labels = load_labels()
        write("hpk_golden.json", build_golden(db, labels))
        write("hpk_prefilter_sample.json", build_prefilter_sample(db))
        write("hpk_admission_sample.json", build_admission_sample(db))
        write("hpk_funnel.json", build_funnel(db))
    finally:
        db.close()


if __name__ == "__main__":
    main()
